"""Formatting helpers for the DBARS defense document."""
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

NAVY = RGBColor(0x1F, 0x35, 0x64)
ACCENT = RGBColor(0x00, 0x7A, 0x6E)
GREY = RGBColor(0x55, 0x5F, 0x6D)
CODE_BG = "F2F4F7"
CALLOUT_BG = "EAF3F1"
WARN_BG = "FDF3E7"


def shade(cell_or_para, hexcolor):
    el = cell_or_para._tc if hasattr(cell_or_para, "_tc") else cell_or_para._p
    pr = el.get_or_add_tcPr() if hasattr(el, "get_or_add_tcPr") else el.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hexcolor)
    pr.append(shd)


def new_document():
    doc = Document()
    for name, size in (("Normal", 10.5),):
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.line_spacing = 1.12
    for lvl, (size, color) in enumerate(
        [(20, NAVY), (15, NAVY), (12.5, ACCENT), (11, NAVY)], start=1
    ):
        st = doc.styles["Heading %d" % lvl]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = color
        st.font.bold = True
        st.paragraph_format.space_before = Pt(14 if lvl < 3 else 10)
        st.paragraph_format.space_after = Pt(6)
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.85)
    sec.top_margin = sec.bottom_margin = Inches(0.8)
    return doc


def h(doc, level, text):
    p = doc.add_heading(text, level=level)
    return p


def para(doc, text, bold=False, italic=False, size=None, color=None, space_after=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    return p


def bullets(doc, items, style="List Bullet"):
    for it in items:
        p = doc.add_paragraph(style=style)
        if isinstance(it, tuple):
            r = p.add_run(it[0]); r.bold = True
            p.add_run(" " + it[1])
        else:
            p.add_run(it)
    return doc


def numbered(doc, items):
    return bullets(doc, items, style="List Number")


def code(doc, text, caption=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Inches(0.12)
    for i, line in enumerate(text.rstrip("\n").split("\n")):
        r = p.add_run(("" if i == 0 else "\n") + line)
        r.font.name = "Consolas"
        r.font.size = Pt(8.6)
    shade(p, CODE_BG)
    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    for side in ("left",):
        b = OxmlElement("w:%s" % side)
        b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "18")
        b.set(qn("w:space"), "6"); b.set(qn("w:color"), "007A6E")
        borders.append(b)
    pPr.append(borders)
    if caption:
        cp = doc.add_paragraph()
        cr = cp.add_run(caption)
        cr.italic = True; cr.font.size = Pt(8.5); cr.font.color.rgb = GREY
        cp.paragraph_format.space_after = Pt(10)
    return p


def callout(doc, title, text, warn=False):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0)
    shade(c, WARN_BG if warn else CALLOUT_BG)
    p0 = c.paragraphs[0]
    r = p0.add_run(title)
    r.bold = True; r.font.size = Pt(10); r.font.color.rgb = NAVY
    p1 = c.add_paragraph()
    r1 = p1.add_run(text); r1.font.size = Pt(9.6)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def table(doc, headers, rows, widths=None, font=8.6):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, htxt in enumerate(headers):
        hdr[i].text = ""
        r = hdr[i].paragraphs[0].add_run(htxt)
        r.bold = True; r.font.size = Pt(font); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(hdr[i], "1F3564")
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            for j, line in enumerate(str(val).split("\n")):
                rr = p.add_run(("" if j == 0 else "\n") + line)
                rr.font.size = Pt(font)
                if str(val).startswith("`"):
                    rr.font.name = "Consolas"
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def caption(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = True; r.font.size = Pt(8.5); r.font.color.rgb = GREY
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    return p


def evidence(doc, text):
    p = doc.add_paragraph()
    r = p.add_run("Evidence: ")
    r.bold = True; r.font.size = Pt(8.6); r.font.color.rgb = ACCENT
    r2 = p.add_run(text)
    r2.font.size = Pt(8.6); r2.font.name = "Consolas"; r2.font.color.rgb = GREY
    p.paragraph_format.space_after = Pt(8)
    return p


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_page_numbers(doc):
    for section in doc.sections:
        footer = section.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("DBARS — Project Defense & Codebase Mastery Document    |    Page ")
        r.font.size = Pt(8); r.font.color.rgb = GREY
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        p._p.append(fld)


def add_header(doc, text):
    for section in doc.sections:
        hp = section.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = hp.add_run(text)
        r.font.size = Pt(8); r.font.color.rgb = GREY; r.italic = True


def add_toc(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    fld = OxmlElement("w:fldChar"); fld.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = "TOC " + chr(92) + "o \"1-2\" " + chr(92) + "h " + chr(92) + "z " + chr(92) + "u"
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t")
    txt.text = "Right-click here and choose 'Update Field' to build the Table of Contents."
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    for el in (fld, instr, sep, txt, end):
        run._r.append(el)
